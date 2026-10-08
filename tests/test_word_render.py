"""Optional desktop Word integration test; never run against a user's document."""
import hashlib
import json
import os
from pathlib import Path
import subprocess
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]


@unittest.skipUnless(os.name == 'nt' and os.environ.get('SKILL_TEST_WORD') == '1',
                     'Opt in on Windows with desktop Word: SKILL_TEST_WORD=1')
class WordRenderTests(unittest.TestCase):
    def test_export_map_source_preservation_and_existing_output_refusal(self):
        from docx import Document
        from docx.enum.text import WD_LINE_SPACING
        from docx.oxml import OxmlElement
        from docx.shared import Pt

        with tempfile.TemporaryDirectory(prefix='scan-skill-test-') as directory:
            folder = Path(directory)
            source, pdf, page_map = folder / 'source.docx', folder / 'render.pdf', folder / 'map.json'
            document = Document()
            document.add_heading('Synthetic renderer regression', level=1)
            p = document.add_paragraph('An editable fraction: ')
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
            p.paragraph_format.line_spacing = Pt(15)
            math, fraction = OxmlElement('m:oMath'), OxmlElement('m:f')
            for tag, value in [('m:num','1'), ('m:den','2')]:
                part, run, text = OxmlElement(tag), OxmlElement('m:r'), OxmlElement('m:t')
                text.text = value
                run.append(text); part.append(run); fraction.append(part)
            math.append(fraction); p._p.append(math)
            document.sections[0].footer.paragraphs[0].text = 'P.1 of 1'
            document.save(source)
            original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
            command = ['powershell', '-NoProfile', '-File', str(ROOT / 'scripts/render_word.ps1'),
                       '-Docx', str(source), '-Pdf', str(pdf), '-MapJson', str(page_map)]
            result = subprocess.run(command, capture_output=True, text=True, timeout=60)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), original_hash)
            self.assertTrue(pdf.read_bytes().startswith(b'%PDF-'))
            report = json.loads(page_map.read_text(encoding='utf-8'))
            self.assertEqual(report['renderer'], 'Microsoft Word')
            self.assertEqual(report['physical_page_count'], 1)
            self.assertTrue(report['separate_instance_guard'])
            pdf_hash = hashlib.sha256(pdf.read_bytes()).hexdigest()
            map_hash = hashlib.sha256(page_map.read_bytes()).hexdigest()
            refused = subprocess.run(command, capture_output=True, text=True, timeout=30)
            self.assertNotEqual(refused.returncode, 0)
            self.assertEqual(hashlib.sha256(pdf.read_bytes()).hexdigest(), pdf_hash)
            self.assertEqual(hashlib.sha256(page_map.read_bytes()).hexdigest(), map_hash)
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), original_hash)


if __name__ == '__main__':
    unittest.main()
